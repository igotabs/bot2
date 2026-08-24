"""
Runtime UI-text loader.

All user-facing strings (messages, button labels, prompts) live OUTSIDE the code
in an external key/value file:
    * Column A -> programmatic KEY   (used in code)
    * Column B -> human-readable VALUE (loaded here at runtime only)

Supported formats: .xlsx (openpyxl) and .csv.
The file path is resolved from the CONFIG_FILE env var, otherwise the first of a
few conventional names that actually exists on disk is used.

The code NEVER hard-codes real texts: every call site uses
    CONFIG.get("SOME_KEY", "safe fallback")
so the bot keeps working even if a key (or the whole file) is missing.
"""

from __future__ import annotations

import csv
import logging
import os
from pathlib import Path
from typing import Dict, Iterable, Optional, Tuple

logger = logging.getLogger(__name__)

# Rows whose key equals one of these (case-insensitive) are treated as headers.
_HEADER_KEYS = {"column1", "key", "ключ", "keys"}

# Candidate file names, tried in order when CONFIG_FILE is not set.
_DEFAULT_CANDIDATES = ("secrets.xlsx", "secret.xlsx", "secrets.csv", "secret.csv")


def _resolve_path() -> Optional[Path]:
    """Return the config file path from env or the first existing candidate."""
    env_path = os.getenv("CONFIG_FILE")
    if env_path:
        p = Path(env_path)
        if p.exists():
            return p
        logger.warning("CONFIG_FILE=%s does not exist; falling back to defaults.", env_path)

    for name in _DEFAULT_CANDIDATES:
        p = Path(name)
        if p.exists():
            return p
    return None


def _iter_rows_xlsx(path: Path) -> Iterable[Tuple[object, object]]:
    """Yield (col_a, col_b) tuples from an .xlsx file."""
    from openpyxl import load_workbook  # imported lazily so .csv-only setups work

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
            key = row[0] if len(row) > 0 else None
            value = row[1] if len(row) > 1 else None
            yield key, value
    finally:
        wb.close()


def _iter_rows_csv(path: Path) -> Iterable[Tuple[object, object]]:
    """Yield (col_a, col_b) tuples from a .csv file (utf-8, comma-separated)."""
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            key = row[0] if len(row) > 0 else None
            value = row[1] if len(row) > 1 else None
            yield key, value


def _load(path: Optional[Path]) -> Dict[str, str]:
    """Parse the config file into a {KEY: VALUE} dict. Never raises."""
    data: Dict[str, str] = {}
    if path is None:
        logger.warning(
            "No config file found (looked for CONFIG_FILE and %s). "
            "Using in-code fallbacks for all UI texts.",
            ", ".join(_DEFAULT_CANDIDATES),
        )
        return data

    try:
        suffix = path.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            rows = _iter_rows_xlsx(path)
        elif suffix == ".csv":
            rows = _iter_rows_csv(path)
        else:
            logger.error("Unsupported config format '%s' for %s.", suffix, path)
            return data

        for key, value in rows:
            if key is None:
                continue
            key_str = str(key).strip()
            if not key_str or key_str.lower() in _HEADER_KEYS:
                continue
            data[key_str] = "" if value is None else str(value)

        logger.info("Loaded %d UI keys from %s.", len(data), path)
    except Exception as exc:  # never let config problems crash the bot
        logger.error("Failed to load config from %s: %s", path, exc)

    return data


class ConfigStore:
    """Dict-like read-only store with a ``.get(key, fallback)`` interface."""

    def __init__(self) -> None:
        self._path: Optional[Path] = _resolve_path()
        self._data: Dict[str, str] = _load(self._path)

    def get(self, key: str, default: str = "") -> str:
        value = self._data.get(key)
        # Treat empty strings as "not configured" so the fallback is used.
        return value if value not in (None, "") else default

    def reload(self) -> "ConfigStore":
        """Re-read the source file at runtime (e.g. after editing texts)."""
        self._path = _resolve_path()
        self._data = _load(self._path)
        return self

    def __contains__(self, key: str) -> bool:
        return key in self._data

    def __len__(self) -> int:
        return len(self._data)


# Singleton imported across the project: `from utils.config_loader import CONFIG`
CONFIG = ConfigStore()
